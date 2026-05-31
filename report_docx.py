"""
Генератор протокола аттестации ФСТЭК (DOCX).
Поддерживает покомпонентный и системный аудит.
"""
import json
from datetime import datetime
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH

class FSTEKReportGenerator:
    def __init__(self, audit_json_path: str):
        with open(audit_json_path, 'r', encoding='utf-8') as f:
            self.data = json.load(f)
        self._normalize()
        self.doc = Document()
        style = self.doc.styles['Normal']
        style.font.name = 'Times New Roman'
        style.font.size = Pt(12)
        style.paragraph_format.line_spacing = 1.5

    def _normalize(self):
        d = self.data
        # Извлекаем classification
        if "classification" in d:
            self.classification = d["classification"]
        else:
            self.classification = d.get("system", {}).get("classification", {})

        # Извлекаем/строим results
        if "requirement_analysis" in d:
            ra = d["requirement_analysis"]
            self.results = []
            req_doc_map = {}
            for r in d.get("requirements", []):
                req_doc_map[r.get("paragraph_id", "")] = r.get("source_document", "")

            for req_id, analysis in ra.items():
                self.results.append({
                    "requirement_id": req_id,
                    "source_document": analysis.get("source_document", "") or req_doc_map.get(req_id, ""),
                    "component_name": ", ".join(analysis.get("responsible_components", ["ПАК"])),
                    "requirement_text": analysis.get("analysis", ""),
                    "compliance_status": analysis.get("status", "UNKNOWN"),
                    "analysis": {
                        "current_state": analysis.get("analysis", ""),
                        "gap": analysis.get("gaps", ""),
                        "risk_level": "HIGH" if analysis.get("status") == "NON_COMPLIANT" else ("MEDIUM" if analysis.get("status") == "PARTIALLY_COMPLIANT" else "LOW"),
                    },
                    "developer_note": analysis.get("recommendation", ""),
                    "recommended_configuration_change": {
                        "new_value": analysis.get("recommendation", ""),
                        "justification": analysis.get("compensating_measures", ""),
                    },
                })
        elif "results" in d:
            self.results = d["results"]
        else:
            self.results = []

        # Извлекаем сводку
        if "overall_assessment" in d:
            oa = d["overall_assessment"]
            self.summary = {
                "total_requirements": oa.get("total_requirements", len(self.results)),
                "compliant": oa.get("fully_compliant", sum(1 for r in self.results if r.get("compliance_status") == "COMPLIANT")),
                "non_compliant": oa.get("non_compliant", sum(1 for r in self.results if r.get("compliance_status") in ["NON_COMPLIANT", "PARTIALLY_COMPLIANT"])),
                "need_more_info": sum(1 for r in self.results if r.get("compliance_status") == "NEED_MORE_INFO"),
            }
        elif "audit_summary" in d:
            self.summary = d["audit_summary"]
        else:
            self.summary = {"total_requirements": len(self.results), "compliant": 0, "non_compliant": 0, "need_more_info": 0}

        # Компоненты
        self.components = d.get("system", {}).get("components", [])
        if not self.components:
            self.components = d.get("components", [])

        # Система
        self.system_name = d.get("system_name", d.get("system", {}).get("system_name", "ПАК"))

    def _h(self, text, level=1):
        h = self.doc.add_heading(text, level=level)
        for r in h.runs:
            r.font.name = 'Times New Roman'

    def _t(self, headers, rows):
        t = self.doc.add_table(rows=1+len(rows), cols=len(headers))
        t.style = 'Table Grid'
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(headers):
            c = t.rows[0].cells[i]; c.text = h
            for p in c.paragraphs:
                for r in p.runs: r.font.bold = True; r.font.size = Pt(9)
        for ri, row in enumerate(rows):
            for ci, val in enumerate(row):
                c = t.rows[ri+1].cells[ci]; c.text = str(val) if val else ""
                for p in c.paragraphs:
                    for r in p.runs: r.font.size = Pt(9)
        return t

    def generate(self, output_path=None):
        # Титул
        self._h('ПРОТОКОЛ АТТЕСТАЦИОННЫХ ИСПЫТАНИЙ', 0)
        self.doc.add_paragraph()

        # Раздел 1: Общие сведения
        self._h('1. Общие сведения')
        self._t(["Параметр", "Значение"], [
            ["Наименование системы", self.system_name],
            ["Дата аудита", datetime.now().strftime("%d.%m.%Y")],
            ["Тип системы", self.classification.get("system_type", "")],
            ["Класс (приказ №117)", self.classification.get("fstek_117_class", "")],
            ["Класс (приказ №21)", self.classification.get("fstek_21_class", "")],
            ["Наличие ДСП", "Да" if self.classification.get("has_dsp") else "Нет"],
            ["Категория ПДн", self.classification.get("ispdn_category", "")],
            ["Субъектов ПДн", str(self.classification.get("ispdn_subjects_count", ""))],
            ["Тип угроз", self.classification.get("threat_type", "")],
            ["Уровень (ПП №1119)", self.classification.get("pp1119_level", "")],
        ])
        self.doc.add_paragraph()

        # Раздел 2: Компоненты
        self._h('2. Состав ПАК')
        if self.components:
            self._t(["№", "Наименование", "Тип"],
                    [[str(i), c.get("component_name", c.get("name", "")), c.get("component_type", c.get("type", ""))]
                     for i, c in enumerate(self.components, 1)])
        else:
            self.doc.add_paragraph("Данные о компонентах не загружены.")
        self.doc.add_paragraph()

        # Раздел 3: Сводка
        self._h('3. Сводная ведомость')
        self._t(["Показатель", "Количество"], [
            ["Всего требований", str(self.summary.get("total_requirements", 0))],
            ["Соответствует", str(self.summary.get("compliant", 0))],
            ["Не соответствует", str(self.summary.get("non_compliant", 0))],
            ["Требует уточнения", str(self.summary.get("need_more_info", 0))],
        ])
        self.doc.add_paragraph()

        # Раздел 4: Матрица соответствия
        self._h('4. Матрица соответствия')
        if self.results:
            self._t(["№", "Документ", "Пункт", "Компонент", "Требование", "Статус", "Рекомендация"],
                    [[str(i), r.get("source_document", ""), r.get("requirement_id", ""),
                      r.get("component_name", ""), (r.get("requirement_text", "") or "")[:200],
                      r.get("compliance_status", ""), r.get("developer_note", "")]
                     for i, r in enumerate(self.results, 1)])
        self.doc.add_paragraph()

        # Раздел 5: Несоответствия
        gaps = [r for r in self.results if r.get("compliance_status") in ["NON_COMPLIANT", "PARTIALLY_COMPLIANT"]]
        self._h('5. Выявленные несоответствия')
        if gaps:
            self._t(["№", "Документ", "Пункт", "Недостаток", "Рекомендация"],
                    [[str(i), r.get("source_document", ""), r.get("requirement_id", ""),
                      r.get("analysis", {}).get("gap", ""), r.get("developer_note", "")]
                     for i, r in enumerate(gaps, 1)])
        else:
            self.doc.add_paragraph("Несоответствий не выявлено.")
        self.doc.add_paragraph()

        # Раздел 6: Заключение
        self._h('6. Заключение')
        total = self.summary.get("total_requirements", 1) or 1
        compliant = self.summary.get("compliant", 0)
        pct = round(compliant / total * 100)
        self.doc.add_paragraph(
            f"По результатам системного аудита ПАК «{self.system_name}» установлено:\n\n"
            f"• Всего требований НПА: {total}\n"
            f"• Соответствует: {compliant} ({pct}%)\n"
            f"• Не соответствует: {self.summary.get('non_compliant', 0)}\n\n"
            f"ВНИМАНИЕ: Протокол сформирован автоматизированной системой. "
            f"Окончательное решение принимается аттестационной комиссией."
        )

        if output_path is None:
            output_path = f"reports/protocol_{self.system_name}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(output_path)
        print(f"📄 Протокол DOCX: {output_path}")
        return output_path

if __name__ == "__main__":
    import sys
    if len(sys.argv) > 1:
        FSTEKReportGenerator(sys.argv[1]).generate()
