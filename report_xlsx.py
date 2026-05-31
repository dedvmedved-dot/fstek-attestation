"""Генератор матрицы соответствия в Excel с поддержкой системного аудита."""
import json
from datetime import datetime
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

class FSTEKExcelGenerator:
    GREEN = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    RED = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    YELLOW = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
    HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
    BORDER = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

    def __init__(self, path):
        with open(path, 'r', encoding='utf-8') as f:
            self.data = json.load(f)
        self._normalize()
        self.wb = Workbook()

    def _normalize(self):
        d = self.data
        self.classification = d.get("classification", d.get("system", {}).get("classification", {}))
        self.system_name = d.get("system_name", d.get("system", {}).get("system_name", "ПАК"))
        self.components = d.get("system", {}).get("components", d.get("components", []))

        if "requirement_analysis" in d:
            ra = d["requirement_analysis"]
            # Маппинг для заполнения source_document
            req_doc_map = {}
            for r in d.get("requirements", []):
                req_doc_map[r.get("paragraph_id", "")] = r.get("source_document", "")
            self.results = []
            for req_id, analysis in ra.items():
                self.results.append({
                    "requirement_id": req_id,
                    "source_document": analysis.get("source_document", "") or req_doc_map.get(req_id, ""),
                    "component_name": ", ".join(analysis.get("responsible_components", ["ПАК"])),
                    "requirement_text": analysis.get("analysis", ""),
                    "compliance_status": analysis.get("status", "UNKNOWN"),
                    "analysis": {"current_state": analysis.get("analysis", ""), "gap": analysis.get("gaps", ""),
                                 "risk_level": "HIGH" if analysis.get("status") == "NON_COMPLIANT" else ("MEDIUM" if analysis.get("status") == "PARTIALLY_COMPLIANT" else "LOW")},
                    "developer_note": analysis.get("recommendation", ""),
                })
        else:
            self.results = d.get("results", [])

        if "overall_assessment" in d:
            oa = d["overall_assessment"]
            self.summary = {"total_requirements": oa.get("total_requirements", len(self.results)),
                           "compliant": oa.get("fully_compliant", 0), "non_compliant": oa.get("non_compliant", 0), "need_more_info": 0}
        else:
            self.summary = d.get("audit_summary", {"total_requirements": len(self.results), "compliant": 0, "non_compliant": 0, "need_more_info": 0})

    def _style_header(self, ws, row, n):
        for c in range(1, n+1):
            cell = ws.cell(row=row, column=c)
            cell.fill = self.HEADER_FILL; cell.font = self.HEADER_FONT
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = self.BORDER

    def _style(self, cell, status=None):
        cell.border = self.BORDER; cell.alignment = Alignment(vertical='top', wrap_text=True)
        if status == "COMPLIANT": cell.fill = self.GREEN
        elif status in ["NON_COMPLIANT", "PARTIALLY_COMPLIANT"]: cell.fill = self.RED
        elif status == "NEED_MORE_INFO": cell.fill = self.YELLOW

    def generate(self, output_path=None):
        ws = self.wb.active; ws.title = "Сводка"
        ws.merge_cells('A1:B1'); ws['A1'] = "СВОДКА РЕЗУЛЬТАТОВ АУДИТА ФСТЭК"; ws['A1'].font = Font(bold=True, size=14)
        rows = [["Система", self.system_name], ["Дата", datetime.now().strftime("%d.%m.%Y")],
                ["Класс", self.classification.get("fstek_117_class", "")], ["Уровень ПДн", self.classification.get("pp1119_level", "")],
                ["Тип угроз", self.classification.get("threat_type", "")], ["ДСП", "Да" if self.classification.get("has_dsp") else "Нет"],
                ["Всего требований", self.summary.get("total_requirements", 0)],
                ["Соответствует", self.summary.get("compliant", 0)], ["Не соответствует", self.summary.get("non_compliant", 0)]]
        for i, (l, v) in enumerate(rows, 3):
            ws.cell(row=i, column=1, value=l).font = Font(bold=True); ws.cell(row=i, column=2, value=v)
        ws.column_dimensions['A'].width = 25; ws.column_dimensions['B'].width = 40

        ws2 = self.wb.create_sheet("Матрица соответствия")
        h = ["№", "Документ", "Пункт НПА", "Компонент", "Требование", "Статус", "Текущее состояние", "Недостаток", "Риск", "Рекомендация"]
        for c, hdr in enumerate(h, 1): ws2.cell(row=1, column=c, value=hdr)
        self._style_header(ws2, 1, len(h))
        for ri, r in enumerate(self.results, 2):
            status = r.get("compliance_status", "")
            a = r.get("analysis", {})
            data = [ri-1, r.get("source_document", ""), r.get("requirement_id", ""), r.get("component_name", ""),
                    (r.get("requirement_text", "") or "")[:900], status, a.get("current_state", ""), a.get("gap", ""),
                    a.get("risk_level", ""), r.get("developer_note", "")]
            for ci, v in enumerate(data, 1):
                cell = ws2.cell(row=ri, column=ci, value=v); self._style(cell, status if ci == 6 else None)
        widths = [5, 18, 14, 22, 45, 14, 30, 30, 8, 30]
        for i, w in enumerate(widths, 1): ws2.column_dimensions[get_column_letter(i)].width = w
        ws2.auto_filter.ref = f"A1:{get_column_letter(len(h))}{len(self.results)+1}"

        ws3 = self.wb.create_sheet("Компоненты")
        ch = ["№", "Наименование", "Тип", "Сертификат"]
        for c, hdr in enumerate(ch, 1): ws3.cell(row=1, column=c, value=hdr)
        self._style_header(ws3, 1, len(ch))
        for ri, comp in enumerate(self.components, 2):
            for ci, v in enumerate([ri-1, comp.get("component_name", comp.get("name", "")),
                                     comp.get("component_type", comp.get("type", "")), comp.get("certificate", "")], 1):
                cell = ws3.cell(row=ri, column=ci, value=v); cell.border = self.BORDER
        ws3.column_dimensions['A'].width = 5; ws3.column_dimensions['B'].width = 40
        ws3.column_dimensions['C'].width = 25; ws3.column_dimensions['D'].width = 30

        if output_path is None:
            output_path = f"reports/matrix_{self.system_name}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        self.wb.save(output_path)
        print(f"📊 Матрица XLSX: {output_path}")
        return output_path

if __name__ == "__main__":
    import sys
    if len(sys.argv) > 1: FSTEKExcelGenerator(sys.argv[1]).generate()
